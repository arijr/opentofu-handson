#!/usr/bin/env python3
"""Build the OpenTofu workshop presentation as a PowerPoint deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "apresentacao.pptx"

NAVY = RGBColor(11, 18, 32)
PANEL = RGBColor(19, 31, 50)
PANEL_ALT = RGBColor(23, 40, 62)
WHITE = RGBColor(244, 248, 252)
MUTED = RGBColor(168, 184, 202)
TEAL = RGBColor(58, 220, 194)
BLUE = RGBColor(95, 163, 255)
GOLD = RGBColor(255, 197, 91)
FONT = "Aptos"

SLIDES = [
    {
        "title": "OpenTofu e fundamentos de IaC",
        "time": "1 min",
        "kicker": "MINICURSO • 30 MIN DE CONTEÚDO",
        "kind": "cover",
        "items": [
            ("O que vamos entender", "História, governança e conceitos essenciais de IaC."),
            ("O que vamos construir", "Um fluxo declarativo, revisável e reproduzível."),
            ("Onde vamos praticar", "Floci local: nenhuma infraestrutura AWS real."),
        ],
        "notes": (
            "Abra explicando que a apresentação dura 30 minutos e antecede os "
            "laboratórios. Reforce desde o início que os exercícios usam apenas "
            "o Floci em localhost:4566, não uma conta AWS."
        ),
    },
    {
        "title": "Terraform: uma década de ecossistema",
        "time": "3 min",
        "kicker": "2014 • ORIGEM",
        "items": [
            ("2014", "Terraform é lançado como software open source sob MPL 2.0."),
            ("Uma linguagem comum", "Configurações declarativas descrevem infraestrutura de provedores diferentes."),
            ("Efeito de ecossistema", "Providers, módulos, ferramentas e conhecimento compartilhado crescem ao redor do projeto."),
        ],
        "notes": (
            "Fato histórico: o manifesto OpenTofu situa a publicação do Terraform "
            "em 2014 sob Mozilla Public License 2.0. Apresente o ecossistema como "
            "resultado de anos de adoção, sem sugerir que havia uma única forma "
            "de uso. Fonte: https://opentofu.org/manifesto/"
        ),
    },
    {
        "title": "A mudança de licença",
        "time": "4 min",
        "kicker": "10 AGO 2023 • HASHICORP",
        "items": [
            ("O anúncio", "HashiCorp anuncia a mudança de Terraform e outros produtos core."),
            ("De MPL para BUSL", "Novas versões passam à Business Source License 1.1."),
            ("O que não significa", "Não é uma revogação retroativa das versões publicadas sob MPL."),
        ],
        "notes": (
            "Diferencie fato e interpretação. HashiCorp anunciou a mudança em "
            "10 de agosto de 2023. O anúncio alterou a licença aplicável às "
            "versões futuras conforme os termos publicados; evite dizer que "
            "mudou retroativamente a licença de versões antigas. Fontes: "
            "https://opentofu.org/manifesto/ e "
            "https://opentofu.org/blog/opentofu-announces-fork-of-terraform/"
        ),
    },
    {
        "title": "Preocupações e resposta comunitária",
        "time": "3 min",
        "kicker": "FATO HISTÓRICO ≠ OPINIÃO",
        "items": [
            ("Preocupação citada", "Parte da comunidade questionou previsibilidade e o alcance da concessão de uso."),
            ("Posição do manifesto", "Os autores consideram a BUSL não open source e defendem uma base aberta."),
            ("Leitura responsável", "É uma posição dos autores e apoiadores — não uma decisão judicial nem consenso universal."),
        ],
        "notes": (
            "A preocupação sobre termos potencialmente competitivos está descrita "
            "pelo manifesto e FAQ do OpenTofu. Deixe explícito que isso representa "
            "a perspectiva dos autores e apoiadores; não ofereça aconselhamento "
            "jurídico nem a apresente como consenso de toda a comunidade. Fontes: "
            "https://opentofu.org/manifesto/ e https://opentofu.org/faq/"
        ),
    },
    {
        "title": "Do anúncio ao fork público",
        "time": "3 min",
        "kicker": "AGO–SET 2023 • OPEN TOFU",
        "kind": "timeline",
        "items": [
            ("25 ago 2023", "Anúncio público da criação do fork chamado OpenTofu."),
            ("Set 2023", "O repositório do projeto é publicado para colaboração aberta."),
            ("Próximo passo", "Organizar processos, registry e caminho até a primeira versão estável."),
        ],
        "notes": (
            "O anúncio de 25 de agosto diz que o fork foi criado. O artigo de "
            "setembro anuncia que o repositório passou a estar público; naquele "
            "momento, a primeira versão estável ainda aparecia como etapa futura. "
            "Fontes: https://opentofu.org/blog/opentofu-announces-fork-of-terraform/ "
            "e https://opentofu.org/blog/the-opentofu-fork-is-now-available/"
        ),
    },
    {
        "title": "Fundação e primeira versão estável",
        "time": "3 min",
        "kicker": "SET 2023 → JAN 2024",
        "kind": "timeline",
        "items": [
            ("Linux Foundation", "Projeto anunciado sob uma fundação neutra em setembro de 2023."),
            ("OpenTofu 1.6.0", "Primeira versão estável publicada em janeiro de 2024."),
            ("Por que importa", "Governança e releases tornam a iniciativa um projeto utilizável e aberto a contribuições."),
        ],
        "notes": (
            "A Linux Foundation anunciou a formação do OpenTofu em 20 de setembro "
            "de 2023. A primeira versão estável 1.6.0 foi anunciada em janeiro "
            "de 2024. Diga que a estrutura de fundação apoia governança neutra, "
            "não que garante automaticamente todos os resultados futuros. Fontes: "
            "https://www.linuxfoundation.org/press/announcing-opentofu e "
            "https://opentofu.org/blog/opentofu-is-going-ga/"
        ),
    },
    {
        "title": "Governança e posicionamento",
        "time": "3 min",
        "kicker": "PROJETO COMUNITÁRIO",
        "items": [
            ("Aberto", "Desenvolvimento, propostas e discussões acontecem publicamente."),
            ("Comunitário e imparcial", "A proposta é avaliar contribuições pelo valor para o projeto, não por um fornecedor."),
            ("Compatibilidade", "OpenTofu nasce como fork do Terraform e mantém foco em configurações compatíveis."),
        ],
        "notes": (
            "Esses termos refletem o posicionamento apresentado pelo projeto em "
            "seu manifesto e anúncio. Convide os alunos a distinguir objetivos "
            "de governança de garantias verificáveis. A FAQ e documentação atual "
            "ajudam a avaliar como o projeto evoluiu. Fontes: "
            "https://opentofu.org/manifesto/ e https://opentofu.org/faq/"
        ),
    },
    {
        "title": "IaC declarativa: declarar o desejado",
        "time": "3 min",
        "kicker": "CONFIGURAÇÃO → PLANO",
        "kind": "flow",
        "items": [
            ("Declarar", "Escrevemos o estado desejado em HCL, não uma lista manual de chamadas API."),
            ("Comparar", "OpenTofu lê a configuração, o state e o que o provider observa."),
            ("Revisar", "O plan mostra as ações propostas antes de qualquer apply."),
        ],
        "notes": (
            "Compare declarativo e imperativo: em vez de instruir passo a passo "
            "cada chamada, descrevemos o resultado esperado. O plano é uma "
            "proposta, não uma garantia de que o ambiente nunca mude. Fonte: "
            "https://opentofu.org/docs/cli/commands/plan/"
        ),
    },
    {
        "title": "As peças da configuração",
        "time": "2 min",
        "kicker": "HCL • PROVIDER • RESOURCE • MODULE",
        "items": [
            ("Provider", "Plugin que conversa com uma API, como AWS S3."),
            ("Resource", "Objeto que a configuração declara e OpenTofu gerencia."),
            ("Variable / output", "Entradas parametrizam; outputs expõem valores do módulo."),
            ("Módulo", "Agrupa configuração reutilizável com inputs e outputs."),
        ],
        "notes": (
            "Use o exemplo do curso: provider AWS aponta explicitamente ao Floci; "
            "aws_s3_bucket declara o bucket. Inputs e outputs formam a interface "
            "de um módulo. O módulo raiz também é um módulo. Fonte: "
            "https://opentofu.org/docs/"
        ),
    },
    {
        "title": "Referências e dependências",
        "time": "1 min",
        "kicker": "GRAFO DE RECURSOS",
        "kind": "flow",
        "items": [
            ("Entradas", "Variáveis e locals calculam nomes e tags."),
            ("Referência", "O módulo recebe valores e declara o bucket."),
            ("Dependência implícita", "Uma referência entre recursos informa a ordem ao OpenTofu."),
        ],
        "notes": (
            "Desenhe a sequência arquivos HCL -> OpenTofu -> AWS Provider -> "
            "Floci local. Uma referência a atributo ou output cria dependência "
            "implícita. Use depends_on somente quando a relação não puder ser "
            "modelada por referência normal. Reforce localhost:4566: sem AWS real."
        ),
    },
    {
        "title": "State e ciclo de execução",
        "time": "3 min",
        "kicker": "INIT → PLAN → APPLY → DESTROY",
        "items": [
            ("init · fmt · validate", "Inicializa providers e módulos; formata e valida a configuração."),
            ("plan · apply", "Revise as mudanças propostas; aplique somente após confirmação."),
            ("State e drift", "State mapeia recursos; drift é diferença observada entre o gerenciado e a infraestrutura."),
            ("Cuidado", "State local pode conter dados sensíveis; sensitive não criptografa."),
        ],
        "notes": (
            "Explique que o provider consulta APIs para atualizar a visão do state; "
            "drift depende de uma diferença observável, não de editar o state. "
            "State local simplifica este curso, mas não coordena colaboração. "
            "State pode conter segredos; sensitive mascara algumas saídas, não "
            "criptografa. O laboratório opcional do Dia 1 demonstra lock nativo "
            "contra o Floci. Fontes: https://opentofu.org/docs/ e "
            "https://opentofu.org/docs/language/settings/backends/s3/"
        ),
    },
    {
        "title": "Perguntas para levar ao laboratório",
        "time": "1 min",
        "kicker": "DISCUSSÃO • 30 MIN",
        "kind": "questions",
        "items": [
            "Por que previsibilidade de licença importa para quem mantém ferramentas?",
            "Quem deve participar das decisões de um projeto de infraestrutura?",
            "Quando um plan deve ser tratado como revisão de mudança?",
            "O que pode dar errado ao compartilhar state copiando arquivos?",
            "Que diferenças podem existir entre Floci e AWS?",
        ],
        "notes": (
            "Escolha uma ou duas perguntas para uma rodada rápida. Transicione "
            "para a preparação: iniciaremos Floci local e validaremos o endpoint "
            "antes de executar OpenTofu. Fontes completas: "
            "https://opentofu.org/manifesto/ ; "
            "https://opentofu.org/blog/opentofu-announces-fork-of-terraform/ ; "
            "https://opentofu.org/blog/the-opentofu-fork-is-now-available/ ; "
            "https://opentofu.org/faq/ ; https://opentofu.org/docs/ ; "
            "https://www.linuxfoundation.org/press/announcing-opentofu ; "
            "https://opentofu.org/blog/opentofu-is-going-ga/ ; "
            "https://github.com/opentofu/opentofu/releases/tag/v1.6.0"
        ),
    },
]


def add_text(slide, text, x, y, w, h, size, color, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(0)
    frame.margin_right = Pt(0)
    frame.margin_top = Pt(0)
    frame.margin_bottom = Pt(0)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align
    return box


def add_panel(slide, x, y, w, h, fill=PANEL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_base(slide, index, time):
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.13), Inches(7.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    add_text(
        slide,
        f"OPENTOFU  /  FUNDAMENTOS DE IaC                                      {time.upper()}  •  {index:02d}/{len(SLIDES):02d}",
        0.55,
        7.12,
        12.1,
        0.2,
        9,
        MUTED,
    )


def add_header(slide, title, kicker):
    add_text(slide, kicker, 0.72, 0.38, 11.8, 0.28, 11, TEAL, True)
    add_text(slide, title, 0.72, 0.82, 11.8, 0.7, 29, WHITE, True)


def add_item_cards(slide, items):
    if len(items) == 4:
        positions = [(0.75, 2.0), (6.85, 2.0), (0.75, 4.05), (6.85, 4.05)]
        card_w, card_h = 5.72, 1.72
    else:
        positions = [(0.75, 2.0 + i * 1.48) for i in range(len(items))]
        card_w, card_h = 11.85, 1.22
    for i, (item, pos) in enumerate(zip(items, positions)):
        label, body = item
        x, y = pos
        add_panel(slide, x, y, card_w, card_h, PANEL if i % 2 == 0 else PANEL_ALT)
        add_text(slide, label, x + 0.28, y + 0.19, card_w - 0.56, 0.32, 16, TEAL, True)
        add_text(slide, body, x + 0.28, y + 0.61, card_w - 0.56, card_h - 0.74, 15, WHITE)


def add_timeline(slide, items):
    y = 3.08
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(y + 0.15), Inches(10.0), Inches(0.045)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()
    gap = 3.8
    for index, (label, body) in enumerate(items):
        x = 0.92 + index * gap
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x + 0.95), Inches(y), Inches(0.36), Inches(0.36)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = TEAL if index != 1 else GOLD
        dot.line.fill.background()
        add_panel(slide, x, 3.74, 3.35, 1.74, PANEL if index != 1 else PANEL_ALT)
        add_text(slide, label, x + 0.22, 3.96, 2.9, 0.32, 16, TEAL, True)
        add_text(slide, body, x + 0.22, 4.38, 2.9, 0.88, 14, WHITE)


def add_flow(slide, items):
    colors = [PANEL, PANEL_ALT, PANEL]
    for index, (label, body) in enumerate(items):
        x = 0.8 + index * 4.13
        add_panel(slide, x, 2.7, 3.55, 2.25, colors[index])
        add_text(slide, f"0{index + 1}", x + 0.24, 2.92, 0.58, 0.35, 18, GOLD, True)
        add_text(slide, label, x + 0.24, 3.43, 3.05, 0.42, 18, TEAL, True)
        add_text(slide, body, x + 0.24, 3.98, 3.05, 0.8, 15, WHITE)
        if index < len(items) - 1:
            add_text(slide, "→", x + 3.63, 3.52, 0.45, 0.45, 25, TEAL, True, PP_ALIGN.CENTER)


def add_questions(slide, items):
    for index, text in enumerate(items):
        y = 1.92 + index * 0.89
        add_panel(slide, 0.82, y, 11.65, 0.68, PANEL if index % 2 == 0 else PANEL_ALT)
        add_text(slide, f"{index + 1:02d}", 1.08, y + 0.16, 0.48, 0.25, 14, TEAL, True)
        add_text(slide, text, 1.75, y + 0.14, 10.3, 0.36, 16, WHITE)


def build():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "OpenTofu e fundamentos de IaC"
    presentation.core_properties.subject = "Apresentação introdutória de 30 minutos"
    presentation.core_properties.author = "Minicurso OpenTofu com Floci"
    presentation.core_properties.keywords = "OpenTofu, IaC, Terraform, Floci"

    blank = presentation.slide_layouts[6]
    for index, content in enumerate(SLIDES, start=1):
        slide = presentation.slides.add_slide(blank)
        add_base(slide, index, content["time"])
        if content.get("kind") == "cover":
            add_text(slide, content["kicker"], 0.86, 0.9, 11.7, 0.35, 13, TEAL, True)
            add_text(slide, content["title"], 0.86, 1.55, 11.55, 1.2, 38, WHITE, True)
            add_text(
                slide,
                "Da história do fork ao primeiro plano declarativo",
                0.88,
                2.85,
                10.9,
                0.5,
                20,
                MUTED,
            )
            for i, (label, body) in enumerate(content["items"]):
                x = 0.88 + i * 4.08
                add_panel(slide, x, 4.15, 3.62, 1.55, PANEL if i != 1 else PANEL_ALT)
                add_text(slide, label, x + 0.22, 4.42, 3.18, 0.35, 16, TEAL, True)
                add_text(slide, body, x + 0.22, 4.88, 3.18, 0.62, 14, WHITE)
        else:
            add_header(slide, content["title"], content["kicker"])
            kind = content.get("kind")
            if kind == "timeline":
                add_timeline(slide, content["items"])
            elif kind == "flow":
                add_flow(slide, content["items"])
            elif kind == "questions":
                add_questions(slide, content["items"])
            else:
                add_item_cards(slide, content["items"])

        notes = slide.notes_slide.notes_text_frame
        notes.text = content["notes"]

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(f"Generated {OUTPUT} ({len(SLIDES)} slides)")


if __name__ == "__main__":
    build()
